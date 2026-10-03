"""Binary file exporters for FMEA-GPT: MIL-STD-1629A Excel and PDF generation with Evidence Provenance."""
import io
from typing import Dict, Any, List
from ..agents.state import FMEAReport, FailureModeEntry, SupportLevel


def generate_fmea_excel(report: FMEAReport) -> io.BytesIO:
    """Generate a professionally styled MIL-STD-1629A Excel spreadsheet with provenance evidence."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIL-STD-1629A FMEA"

    # Color Palette
    NAVY_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    CRIT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    WARN_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

    WHITE_BOLD = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
    NORMAL_FONT = Font(name="Segoe UI", size=9)
    ALERT_FONT = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )

    # 1. Title Banner
    ws.merge_cells("A1:N1")
    ws["A1"] = "FAILURE MODE AND EFFECTS ANALYSIS (FMECA) WORKSHEET — MIL-STD-1629A REFERENCE METHODOLOGY"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # 2. Metadata Info Block
    comp = report.component
    meta_rows = [
        ("Component Name:", comp.component_name, "Part Number:", comp.part_number or "N/A"),
        ("System / Subsystem:", f"{comp.system} / {comp.subsystem}", "Airworthiness Class:", comp.regulatory_class),
        ("Operating Environment:", comp.operating_environment, "Primary Function:", comp.primary_function)
    ]

    current_row = 3
    for r in meta_rows:
        ws.cell(row=current_row, column=1, value=r[0]).font = BOLD_FONT
        ws.cell(row=current_row, column=2, value=r[1]).font = NORMAL_FONT
        ws.cell(row=current_row, column=6, value=r[2]).font = BOLD_FONT
        ws.cell(row=current_row, column=7, value=r[3]).font = NORMAL_FONT
        current_row += 1

    current_row += 1

    # 3. Table Column Headers
    headers = [
        "Item ID",
        "Failure Mode",
        "Physical Mechanism",
        "Root Cause",
        "Local Effect",
        "Next Higher Effect",
        "End Effect (Aircraft)",
        "Severity Category",
        "Task 102 Criticality",
        "Detection Method",
        "Recommended Action",
        "Inspection Interval",
        "Grounding Status",
        "Legacy RPN (Compatibility)"
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h)
        cell.font = WHITE_BOLD
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[current_row].height = 32
    current_row += 1

    # 4. Data Rows
    for fm in report.failure_modes:
        is_spf = fm.single_point_failure or fm.severity >= 9
        row_fill = CRIT_FILL if is_spf else (WARN_FILL if fm.rpn >= 100 else None)

        # Categorical Severity representation
        sev_cat = (
            fm.severity_classification.category.value
            if fm.severity_classification and fm.severity_classification.category
            else fm.mil_std_severity_category
        )

        # Task 102 Criticality representation
        if fm.criticality_analysis and fm.criticality_analysis.matrix_position:
            crit_display = fm.criticality_analysis.matrix_position
        elif fm.criticality_analysis:
            crit_display = "INSUFFICIENT EVIDENCE - Fleet rate uncataloged in corpus"
        else:
            crit_display = "Unquantified"

        legacy_str = f"RPN={fm.rpn} (S={fm.severity}, O={fm.occurrence}, D={fm.detection})"

        row_data = [
            fm.mode_id,
            fm.failure_mode,
            fm.physical_mechanism or "-",
            fm.root_cause,
            fm.local_effect,
            fm.next_higher_effect,
            fm.end_effect,
            sev_cat,
            crit_display,
            fm.detection_method,
            fm.recommended_action,
            fm.inspection_interval,
            fm.mode_status.value if hasattr(fm.mode_status, "value") else str(fm.mode_status),
            legacy_str
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = ALERT_FONT if (col_idx in [8] and is_spf) else NORMAL_FONT
            if row_fill:
                cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [1, 8, 9, 13, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        ws.row_dimensions[current_row].height = 50
        current_row += 1

    # Auto-adjust column widths for Worksheet 1
    column_widths = {
        1: 12, 2: 28, 3: 20, 4: 28, 5: 24, 6: 24, 7: 28,
        8: 22, 9: 26, 10: 24, 11: 30, 12: 24, 13: 16, 14: 20
    }
    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 5. Worksheet 2: Evidence & Citations Provenance
    ws_evid = wb.create_sheet(title="Evidence & Provenance")
    ws_evid.merge_cells("A1:H1")
    ws_evid["A1"] = "EVIDENCE PROVENANCE & TECHNICAL CITATIONS — MIL-STD-1629A / FAA AC 33.75-1A"
    ws_evid["A1"].font = TITLE_FONT
    ws_evid["A1"].fill = HEADER_FILL
    ws_evid["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_evid.row_dimensions[1].height = 32

    evid_headers = [
        "Mode ID",
        "Failure Mode",
        "Support Level",
        "Document Title",
        "Source File",
        "Page",
        "Official Report No.",
        "Evidence Excerpt"
    ]
    for c_idx, eh in enumerate(evid_headers, start=1):
        cell = ws_evid.cell(row=3, column=c_idx, value=eh)
        cell.font = WHITE_BOLD
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_evid.row_dimensions[3].height = 25

    e_row = 4
    for fm in report.failure_modes:
        records = fm.evidence_records or [c.to_evidence_record() for c in fm.citations]
        if not records:
            ws_evid.cell(row=e_row, column=1, value=fm.mode_id).font = BOLD_FONT
            ws_evid.cell(row=e_row, column=2, value=fm.failure_mode).font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=3, value=fm.mode_status.value if hasattr(fm.mode_status, "value") else str(fm.mode_status)).font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=4, value="Domain Heuristic (No Direct Corpus Chunk)").font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=5, value="N/A").font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=6, value="-").font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=7, value="-").font = NORMAL_FONT
            ws_evid.cell(row=e_row, column=8, value="Grounding established via turbomachinery engineering baseline.").font = NORMAL_FONT
            e_row += 1
        else:
            for er in records:
                supp_val = er.support_level.value if hasattr(er.support_level, "value") else str(er.support_level)
                ws_evid.cell(row=e_row, column=1, value=fm.mode_id).font = BOLD_FONT
                ws_evid.cell(row=e_row, column=2, value=fm.failure_mode).font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=3, value=supp_val).font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=4, value=er.document_title).font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=5, value=er.source_document).font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=6, value=er.page_number).font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=7, value=er.official_report_number or "-").font = NORMAL_FONT
                ws_evid.cell(row=e_row, column=8, value=er.excerpt).font = NORMAL_FONT
                e_row += 1

    evid_widths = {1: 14, 2: 28, 3: 20, 4: 32, 5: 28, 6: 10, 7: 22, 8: 60}
    for col_idx, width in evid_widths.items():
        ws_evid.column_dimensions[get_column_letter(col_idx)].width = width

    # 6. Worksheet 3: Evidence Coverage & Assumptions
    ws_cov = wb.create_sheet(title="Coverage & Scope")
    ws_cov.merge_cells("A1:F1")
    ws_cov["A1"] = "EVIDENCE COVERAGE SUMMARY & ENGINEERING ASSUMPTIONS"
    ws_cov["A1"].font = TITLE_FONT
    ws_cov["A1"].fill = HEADER_FILL
    ws_cov["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_cov.row_dimensions[1].height = 32

    # Coverage breakdown
    cov = report.evidence_coverage
    if cov:
        cov_rows = [
            ("Total Claims Evaluated", cov.total_claims),
            ("Direct Source Claims", cov.directly_supported),
            ("Supporting Source Claims", cov.supporting_evidence),
            ("Engineering Inference Claims", cov.engineering_inference),
            ("Domain Heuristic Claims", cov.domain_heuristic),
            ("Insufficient Evidence Claims", cov.insufficient_evidence),
            ("Unsupported Claims", cov.unsupported),
            ("Evidence Coverage Adequate", "YES" if cov.is_coverage_adequate else "NO (Evidence Gap)")
        ]
        ws_cov.cell(row=3, column=1, value="Metric").font = WHITE_BOLD
        ws_cov.cell(row=3, column=1).fill = NAVY_FILL
        ws_cov.cell(row=3, column=2, value="Count").font = WHITE_BOLD
        ws_cov.cell(row=3, column=2).fill = NAVY_FILL

        for c_idx, (m, v) in enumerate(cov_rows, start=4):
            ws_cov.cell(row=c_idx, column=1, value=m).font = BOLD_FONT
            ws_cov.cell(row=c_idx, column=2, value=str(v)).font = NORMAL_FONT
            ws_cov.cell(row=c_idx, column=1).border = thin_border
            ws_cov.cell(row=c_idx, column=2).border = thin_border

    # Scope & Limitations Notice
    notice_row = 14
    ws_cov.cell(row=notice_row, column=1, value="Engineering Scope & Review Notice").font = WHITE_BOLD
    ws_cov.cell(row=notice_row, column=1).fill = SECTION_FILL
    ws_cov.merge_cells(f"A{notice_row}:F{notice_row}")

    notices = [
        "1. This candidate FMECA worksheet is generated using an evidence-grounded reference methodology (MIL-STD-1629A / FAA AC 33.75-1A).",
        "2. All claims are grounded against the indexed technical corpus; ungrounded failure rates and cycle intervals are explicitly suppressed.",
        "3. Quantitative fleet occurrence rates require OEM proprietary service logs; absent empirical data, probability is designated UNKNOWN.",
        "4. This output is an engineering baseline tool for qualified analysts and does NOT constitute certification sign-off or airworthiness approval."
    ]
    for idx, n in enumerate(notices, start=notice_row + 1):
        ws_cov.cell(row=idx, column=1, value=n).font = NORMAL_FONT
        ws_cov.merge_cells(f"A{idx}:F{idx}")

    ws_cov.column_dimensions["A"].width = 35
    ws_cov.column_dimensions["B"].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_fmea_pdf(report: FMEAReport) -> io.BytesIO:
    """Generate a clean, high-resolution MIL-STD-1629A PDF document with Provenance Appendix."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=13,
        textColor=colors.HexColor("#0F172A"),
        leading=15,
        spaceAfter=3
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.HexColor("#475569"),
        leading=10,
        spaceAfter=6
    )
    header_style = ParagraphStyle(
        'ColHeader',
        parent=styles['Normal'],
        fontSize=7,
        fontName='Helvetica-Bold',
        textColor=colors.white,
        alignment=1,
        leading=9
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=6.5,
        leading=8,
        textColor=colors.HexColor("#1E293B")
    )
    center_cell_style = ParagraphStyle(
        'CenterCell',
        parent=cell_style,
        alignment=1
    )
    alert_cell_style = ParagraphStyle(
        'AlertCell',
        parent=cell_style,
        alignment=1,
        fontName='Helvetica-Bold',
        textColor=colors.HexColor("#991B1B")
    )

    story = []

    # Title Banner
    story.append(Paragraph("FAILURE MODE AND EFFECTS ANALYSIS (FMECA) REPORT", title_style))
    story.append(Paragraph(
        "MIL-STD-1629A Reference Methodology & FAA AC 33.75-1A Safety Analysis &bull; Requires Qualified Engineer Review",
        subtitle_style
    ))

    # Meta Table
    comp = report.component
    meta_data = [
        [
            Paragraph(f"<b>Component:</b> {comp.component_name}", cell_style),
            Paragraph(f"<b>Part Number:</b> {comp.part_number or 'N/A'}", cell_style),
            Paragraph(f"<b>System / Subsystem:</b> {comp.system} / {comp.subsystem}", cell_style)
        ],
        [
            Paragraph(f"<b>Airworthiness Tier:</b> {comp.regulatory_class}", cell_style),
            Paragraph(f"<b>Indenture Level:</b> {comp.analysis_scope}", cell_style),
            Paragraph(f"<b>Operating Limits:</b> {comp.operating_environment[:65]}...", cell_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[240, 240, 260])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 6))

    # Compact Evidence Coverage Summary
    cov = report.evidence_coverage
    if cov:
        cov_text = (
            f"<b>Evidence Coverage:</b> Total Claims: {cov.total_claims} | "
            f"Direct Source: {cov.directly_supported} | Supporting Source: {cov.supporting_evidence} | "
            f"Engineering Inference: {cov.engineering_inference} | Domain Heuristic: {cov.domain_heuristic} | "
            f"Insufficient Evidence: {cov.insufficient_evidence} | Unsupported: {cov.unsupported}"
        )
        cov_table = Table([[Paragraph(cov_text, cell_style)]], colWidths=[740])
        cov_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        story.append(cov_table)
        story.append(Spacer(1, 6))

    # Main FMEA Table
    col_widths = [45, 95, 80, 75, 80, 80, 80, 115, 90]
    table_headers = [
        Paragraph("ID", header_style),
        Paragraph("Failure Mode & Mechanism", header_style),
        Paragraph("Root Cause", header_style),
        Paragraph("Local Effect", header_style),
        Paragraph("End Aircraft Effect", header_style),
        Paragraph("Severity (MIL-STD)", header_style),
        Paragraph("Task 102 Criticality", header_style),
        Paragraph("Detection & Mitigation Action", header_style),
        Paragraph("Legacy RPN (SAE J1739)", header_style)
    ]

    table_data = [table_headers]

    for fm in report.failure_modes:
        is_spf = fm.single_point_failure or fm.severity >= 9

        sev_cat = (
            fm.severity_classification.category.value.replace("Category ", "Cat ")
            if fm.severity_classification and fm.severity_classification.category
            else fm.mil_std_severity_category.replace("Category ", "Cat ")
        )
        if is_spf:
            sev_cat += " [SPF]"

        crit_pos = (
            fm.criticality_analysis.matrix_position
            if fm.criticality_analysis and fm.criticality_analysis.matrix_position
            else "Unquantified"
        )

        mech_str = f"<br/><i>({fm.physical_mechanism})</i>" if fm.physical_mechanism else ""
        mode_mech = f"<b>{fm.failure_mode}</b>{mech_str}"
        det_act = f"{fm.detection_method}<br/><b>Action:</b> {fm.recommended_action}"
        rpn_cell = f"RPN={fm.rpn}<br/>(S={fm.severity}, O={fm.occurrence}, D={fm.detection})"

        row = [
            Paragraph(f"<b>{fm.mode_id}</b>", cell_style),
            Paragraph(mode_mech, cell_style),
            Paragraph(fm.root_cause, cell_style),
            Paragraph(fm.local_effect, cell_style),
            Paragraph(fm.end_effect, cell_style),
            Paragraph(sev_cat, alert_cell_style if is_spf else cell_style),
            Paragraph(crit_pos, cell_style),
            Paragraph(det_act, cell_style),
            Paragraph(rpn_cell, center_cell_style)
        ]
        table_data.append(row)

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 2.5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 2.5),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]

    for idx, fm in enumerate(report.failure_modes, start=1):
        if fm.single_point_failure or fm.severity >= 9:
            t_style.append(('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#FEE2E2")))
        elif idx % 2 == 0:
            t_style.append(('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#F8FAFC")))

    main_table.setStyle(TableStyle(t_style))
    story.append(main_table)

    # 6. Provenance & Citations Appendix Table
    story.append(Spacer(1, 8))
    story.append(Paragraph("<b>Airworthiness Evidence & Regulatory Provenance Appendix</b>", subtitle_style))
    evid_headers = [
        Paragraph("Mode ID", header_style),
        Paragraph("Support Level", header_style),
        Paragraph("Document Title & Official ID", header_style),
        Paragraph("Source File", header_style),
        Paragraph("Page", header_style),
        Paragraph("Grounding Excerpt", header_style)
    ]
    evid_data = [evid_headers]

    for fm in report.failure_modes:
        records = fm.evidence_records or [c.to_evidence_record() for c in fm.citations]
        if not records:
            evid_data.append([
                Paragraph(f"<b>{fm.mode_id}</b>", cell_style),
                Paragraph("DOMAIN_HEURISTIC", cell_style),
                Paragraph("Turbomachinery Engineering Baseline", cell_style),
                Paragraph("-", center_cell_style),
                Paragraph("-", center_cell_style),
                Paragraph("Heuristic degradation baseline; OEM data required for certification.", cell_style)
            ])
        else:
            for er in records[:2]:  # Top 2 evidence records per mode to fit page
                supp_val = er.support_level.value if hasattr(er.support_level, "value") else str(er.support_level)
                doc_title = er.document_title
                if er.official_report_number:
                    doc_title += f" [{er.official_report_number}]"
                evid_data.append([
                    Paragraph(f"<b>{fm.mode_id}</b>", cell_style),
                    Paragraph(supp_val.upper(), cell_style),
                    Paragraph(doc_title, cell_style),
                    Paragraph(er.source_document, cell_style),
                    Paragraph(str(er.page_number), center_cell_style),
                    Paragraph(er.excerpt[:160] + "...", cell_style)
                ])

    evid_table = Table(evid_data, colWidths=[45, 90, 180, 130, 30, 265], repeatRows=1)
    evid_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    story.append(evid_table)

    # 7. Limitations & Engineering Review Notice
    story.append(Spacer(1, 6))
    notice_text = (
        "<b>Limitations & Review Notice:</b> This automated FMECA analysis employs the MIL-STD-1629A reference methodology. "
        "Quantitative fleet failure occurrence rates and life intervals require proprietary OEM engine service data; where absent in the public corpus, "
        "occurrence rates are uncataloged. RPN is an automotive metric (SAE J1739) provided solely for legacy comparison. "
        "This document is a candidate analysis requiring review by a qualified aerospace systems engineer and does not constitute formal airworthiness certification."
    )
    notice_table = Table([[Paragraph(notice_text, cell_style)]], colWidths=[740])
    notice_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(notice_table)

    doc.build(story)
    output.seek(0)
    return output
